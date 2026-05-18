#!/usr/bin/env python3
"""توليد ملف Excel لخريطة حقول استمارة 50"""
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

wb = Workbook()

# ══════════════════════════════════════════════════════════════
# البيانات
# ══════════════════════════════════════════════════════════════
SECTIONS = [
    {
        'title': 'أعلى الصفحة — الصندوق الأيمن',
        'color': '1F4E79',
        'rows': [
            (1,  'رقم المسلسل',              '55/2024-2025/0001',              '90.06', '8.97',  False, False),
            (2,  'تاريخ ورود الاستمارة',     '15/11/2024',                     '90.05', '10.79', True,  False),
        ]
    },
    {
        'title': 'قسم أ — اللوح الأيسر (المصلحة والطلب)',
        'color': '1F4E79',
        'rows': [
            (3,  'اسم المصلحة',               'مصلحة الشؤون المالية',           '23.84', '9.30',  False, False),
            (4,  'القسم / الإدارة',           'إدارة المشتريات والمخازن',       '23.84', '11.02', False, False),
            (5,  'المبلغ المستحق إلى',        'شركة الاختبار والتطوير',         '19.84', '12.84', False, False),
            (6,  'رقم الارتباط / بموجب',      'CMT-2024-112',                   '19.84', '16.36', False, False),
            (7,  'صار مراجعته',               'شركة الاختبار والتطوير',         '5.85',  '19.69', False, False),
            (8,  'إذن صرف على البنك',         'بنك مصر',                        '18.87', '21.52', False, False),
            (9,  'شيك على الشارج',            'شركة الاختبار والتطوير',         '18.87', '25.04', False, False),
            (10, 'يسحب باسم',                 'شركة الاختبار والتطوير',         '18.88', '28.25', False, False),
            (11, 'ويرسل إليه على العنوان',    'عنوان المورد',                   '23.89', '29.84', False, False),
        ]
    },
    {
        'title': 'جدول الفواتير — اللوح الأيمن',
        'color': '375623',
        'rows': [
            (12, 'فاتورة 1 — رقم',            '0801',   '59.74', '15.24', False, False),
            (13, 'فاتورة 1 — تاريخ',          '05-10',  '64.41', '15.24', True,  False),
            (14, 'فاتورة 1 — جنيه',           '18500',  '74.47', '15.24', False, True),
            (15, 'فاتورة 1 — قرش',            '00',     '80.74', '15.24', False, False),
            (16, 'فاتورة 2 — رقم',            '0802',   '59.74', '19.44', False, False),
            (17, 'فاتورة 2 — تاريخ',          '12-10',  '64.41', '19.44', True,  False),
            (18, 'فاتورة 2 — جنيه',           '22000',  '74.47', '19.44', False, True),
            (19, 'فاتورة 2 — قرش',            '00',     '80.75', '19.44', False, False),
            (20, 'فاتورة 3 — رقم',            '0803',   '59.73', '22.51', False, False),
            (21, 'فاتورة 3 — تاريخ',          '20-10',  '64.39', '22.51', True,  False),
            (22, 'فاتورة 3 — جنيه',           '7910',   '74.46', '22.51', False, True),
            (23, 'فاتورة 3 — قرش',            '00',     '80.74', '22.51', False, False),
            (24, 'فاتورة 4 — رقم',            '(فارغ)', '59.73', '26.29', False, False),
            (25, 'فاتورة 4 — تاريخ',          '(فارغ)', '64.41', '26.29', True,  False),
            (26, 'فاتورة 4 — جنيه',           '(فارغ)', '74.47', '26.29', False, True),
            (27, 'فاتورة 4 — قرش',            '(فارغ)', '80.73', '26.29', False, False),
            (28, 'الجملة — جنيه',             '48410',  '74.47', '29.15', False, True),
            (29, 'الجملة — قرش',              '00',     '80.74', '29.15', False, True),
        ]
    },
    {
        'title': 'أسفل قسم أ',
        'color': '1F4E79',
        'rows': [
            (30, 'الكاتب المنوط',             'محمد أحمد السيد',   '68.39', '31.91', False, False),
            (31, 'تقييد في سجل (ز)',           'ز/2024/0042',       '25.63', '31.91', False, False),
            (32, 'تاريخ الختم (قسم ب)',        '15/11/2024',        '86.59', '35.07', True,  False),
            (33, 'عدد المرفقات',               '7',                 '5.84',  '40.08', False, False),
        ]
    },
    {
        'title': 'جدول تصنيف الموازنة',
        'color': '7B3F00',
        'rows': [
            (34, 'بند',                        'أثاث ومعدات مكتبية', '47.61', '42.82', False, False),
            (35, 'فصل',                        '3',                  '52.34', '42.82', False, False),
            (36, 'فرع / نوع',                  '(من budget_line)',   '57.18', '42.82', False, False),
            (37, 'قسم / باب',                  '2',                  '61.91', '42.81', False, False),
            (38, 'إجمالي — جنيه',              '48410',              '68.44', '42.81', False, True),
            (39, 'إجمالي — قرش',               '00',                 '75.03', '42.81', False, False),
        ]
    },
    {
        'title': 'جدول الاستقطاعات',
        'color': '7B3F00',
        'rows': [
            (40, 'إجمالي الأصل — جنيه',        '48410', '68.44', '46.56', False, False),
            (41, 'إجمالي الأصل — قرش',         '00',    '75.03', '46.56', False, False),
            (42, 'دمغة عادية — جنيه',           '484',   '68.44', '48.71', False, False),
            (43, 'دمغة عادية — قرش',            '10',    '75.03', '48.71', False, False),
            (44, 'دمغة إضافية — جنيه',          '1452',  '68.44', '50.26', False, False),
            (45, 'دمغة إضافية — قرش',           '30',    '75.03', '50.26', False, False),
            (46, 'دمغة نسبية — جنيه',           '387',   '68.44', '51.87', False, False),
            (47, 'دمغة نسبية — قرش',            '00',    '75.03', '51.87', False, False),
            (48, 'ضريبة الأرباح — جنيه',        '1452',  '68.44', '53.57', False, False),
            (49, 'ضريبة الأرباح — قرش',         '30',    '75.04', '53.57', False, False),
            (50, 'صافي القيمة — جنيه',          '44634', '68.44', '55.34', False, True),
            (51, 'صافي القيمة — قرش',           '30',    '75.02', '55.35', False, True),
        ]
    },
    {
        'title': 'التفقيط والتاريخ',
        'color': '1F4E79',
        'rows': [
            (52, 'الصافي بالكلام (التفقيط)',   'أربعة وأربعون ألف وستمئة وأربعة وثلاثون جنيه', '9.19',  '57.29', False, False),
            (53, 'في سنة (إقرار)',              '٥ (رقم السنة بالعربي)',                          '84.01', '58.17', False, False),
            (76, 'في — يوم-شهر',               '15-11',                                           '91.19', '57.89', True,  False),
        ]
    },
    {
        'title': 'توقيعات قسم ب',
        'color': '1F4E79',
        'rows': [
            (55, 'مراقب الحسابات',             'اسم المستخدم',          '6.83',  '46.43', False, False),
            (56, 'رئيس الحسابات',              'اسم المستخدم',          '6.83',  '48.95', False, False),
            (57, 'رقم حساب البنك',             '012345678901234',       '61.44', '65.21', False, False),
            (58, 'بتاريخ',                     '15/11/2024',            '30.94', '73.13', True,  False),
            (59, 'رئيس المصلحة',               'اسم المستخدم',          '6.83',  '73.13', False, False),
        ]
    },
    {
        'title': 'قسم ج — المراجعة',
        'color': '6A0572',
        'rows': [
            (60, 'تاريخ الختم (ج)',            '15/11/2024',                     '87.54', '77.59', True,  False),
            (61, 'قيد في سجل رقم 55',          '55/2024-2025/0001',              '49.93', '70.73', False, False),
            (62, 'روجع في — سنة',              '٥',                              '55.45', '72.21', False, False),
            (63, 'روجع في — تاريخ',            '20-11',                          '63.68', '72.33', True,  False),
            (64, 'شيك — اسم المستفيد',         'شركة الاختبار والتطوير',         '52.92', '83.15', False, False),
            (65, 'يعتمد سحب — المبلغ',         '44634',                          '47.44', '86.13', False, True),
            (66, 'وكيل / مدير الحسابات',       'اسم المراجع',                    '74.03', '89.11', False, False),
            (67, 'رئيس الحسابات',              'اسم رئيس الحسابات',              '27.95', '88.16', False, False),
            (68, 'في سنة (صرف)',               '٥',                              '40.87', '89.65', False, False),
            (69, 'بمبلغ (رقم)',                '44634',                          '25.10', '81.52', False, True),
        ]
    },
    {
        'title': 'قسم د — الدفع',
        'color': '7B0000',
        'rows': [
            (70, 'تاريخ الختم (د)',            '15/11/2024',             '87.54', '94.10', True,  False),
            (71, 'رقم القيد في دفتر 224',      '224/2024/0155',          '36.00', '84.35', False, True),
            (72, 'إمضاء الكاتب المنوط',        'محمد أحمد السيد',        '11.27', '91.09', False, False),
            (73, 'إمضاء موظفي الشطب',          'أحمد محمد علي',          '38.89', '89.87', False, False),
            (74, 'رقم أمر الدفع',              'PO-2024-0387',           '39.35', '91.37', False, False),
            (75, 'سحب / شيك — اسم',            'شركة الاختبار والتطوير', '64.29', '93.04', False, False),
        ]
    },
]

CONVERSION = [
    ('1 مم',  '+0.476%', '+0.337%'),
    ('2 مم',  '+0.952%', '+0.674%'),
    ('3 مم',  '+1.429%', '+1.011%'),
    ('5 مم',  '+2.381%', '+1.685%'),
    ('1 سم',  '+4.762%', '+3.367%'),
    ('2 سم',  '+9.524%', '+6.734%'),
    ('3 سم',  '+14.286%', '+10.101%'),
]

# ══════════════════════════════════════════════════════════════
# ورقة 1: خريطة الحقول
# ══════════════════════════════════════════════════════════════
ws = wb.active
ws.title = 'خريطة الحقول'
ws.sheet_view.rightToLeft = True

thin = Side(style='thin', color='CCCCCC')
thick = Side(style='medium', color='999999')
border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
border_thick = Border(left=thick, right=thick, top=thick, bottom=thick)

# ترويسة
headers = ['رقم الحقل', 'اسم الحقل', 'البيانات التجريبية', 'x%  (أفقي)', 'y%  (رأسي)', 'تاريخ؟', 'بولد؟']
col_widths = [12, 32, 42, 12, 12, 10, 10]

for i, (h, w) in enumerate(zip(headers, col_widths), 1):
    ws.column_dimensions[get_column_letter(i)].width = w

header_fill = PatternFill('solid', fgColor='1F2D3D')
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)

for i, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=i, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = border_thick

ws.row_dimensions[1].height = 28

row = 2
for section in SECTIONS:
    # عنوان القسم
    sec_fill = PatternFill('solid', fgColor=section['color'])
    sec_font = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
    sec_align = Alignment(horizontal='center', vertical='center', readingOrder=2)

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    cell = ws.cell(row=row, column=1, value=section['title'])
    cell.fill = sec_fill
    cell.font = sec_font
    cell.alignment = sec_align
    cell.border = border_thick
    ws.row_dimensions[row].height = 22
    row += 1

    for (num, label, sample, x, y, is_date, is_bold) in section['rows']:
        # رقم الحقل
        c_num = ws.cell(row=row, column=1, value=num)
        c_num.font = Font(name='Calibri', bold=True, size=11,
                          color='0033CC' if is_date else 'CC0000')
        c_num.alignment = Alignment(horizontal='center', vertical='center')
        c_num.fill = PatternFill('solid', fgColor='EEF2FF' if is_date else 'FFF8F8')
        c_num.border = border_thin

        # اسم الحقل
        c_lbl = ws.cell(row=row, column=2, value=label)
        c_lbl.font = Font(name='Calibri', bold=is_bold, size=10,
                          color='0033CC' if is_date else '111111')
        c_lbl.alignment = Alignment(horizontal='right', vertical='center',
                                    wrap_text=True, readingOrder=2)
        c_lbl.fill = PatternFill('solid', fgColor='F0F4FF' if is_date else 'FFFFFF')
        c_lbl.border = border_thin

        # البيانات التجريبية
        c_sample = ws.cell(row=row, column=3, value=sample)
        c_sample.font = Font(name='Courier New', size=9,
                             color='006600' if '(فارغ)' not in sample else '999999')
        c_sample.alignment = Alignment(horizontal='right', vertical='center',
                                       wrap_text=True, readingOrder=2)
        c_sample.border = border_thin

        # x%
        c_x = ws.cell(row=row, column=4, value=float(x))
        c_x.font = Font(name='Courier New', size=10, bold=True, color='1F4E79')
        c_x.alignment = Alignment(horizontal='center', vertical='center')
        c_x.number_format = '0.00"%"'
        c_x.border = border_thin

        # y%
        c_y = ws.cell(row=row, column=5, value=float(y))
        c_y.font = Font(name='Courier New', size=10, bold=True, color='7B3F00')
        c_y.alignment = Alignment(horizontal='center', vertical='center')
        c_y.number_format = '0.00"%"'
        c_y.border = border_thin

        # تاريخ؟
        c_date = ws.cell(row=row, column=6, value='🔵 نعم' if is_date else '')
        c_date.font = Font(name='Calibri', size=10, color='0033CC' if is_date else '999999')
        c_date.alignment = Alignment(horizontal='center', vertical='center')
        c_date.border = border_thin
        if is_date:
            c_date.fill = PatternFill('solid', fgColor='E8F0FE')

        # بولد؟
        c_bold = ws.cell(row=row, column=7, value='✅ بولد' if is_bold else '')
        c_bold.font = Font(name='Calibri', size=10, bold=is_bold,
                           color='006600' if is_bold else '999999')
        c_bold.alignment = Alignment(horizontal='center', vertical='center')
        c_bold.border = border_thin
        if is_bold:
            c_bold.fill = PatternFill('solid', fgColor='F0FFF0')

        ws.row_dimensions[row].height = 20
        row += 1

# ══════════════════════════════════════════════════════════════
# ورقة 2: جدول التحويل
# ══════════════════════════════════════════════════════════════
ws2 = wb.create_sheet('جدول التحويل مم → %')
ws2.sheet_view.rightToLeft = True

ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 18
ws2.column_dimensions['C'].width = 18

conv_headers = ['المسافة', 'أفقي x% (يمين +)', 'رأسي y% (تحت +)']
for i, h in enumerate(conv_headers, 1):
    cell = ws2.cell(row=1, column=i, value=h)
    cell.fill = PatternFill('solid', fgColor='1F2D3D')
    cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    cell.alignment = Alignment(horizontal='center', vertical='center', readingOrder=2)
    cell.border = border_thick
ws2.row_dimensions[1].height = 26

alt_colors = ['FFFFFF', 'F2F7FF']
for i, (dist, xp, yp) in enumerate(CONVERSION, 2):
    fill = PatternFill('solid', fgColor=alt_colors[i % 2])
    for col, val in enumerate([dist, xp, yp], 1):
        cell = ws2.cell(row=i, column=col, value=val)
        cell.font = Font(name='Courier New', size=11,
                         bold=(col == 1),
                         color='1F4E79' if col == 2 else '7B3F00' if col == 3 else '111111')
        cell.alignment = Alignment(horizontal='center', vertical='center', readingOrder=2)
        cell.fill = fill
        cell.border = border_thin
    ws2.row_dimensions[i].height = 22

# ملاحظة
ws2.cell(row=len(CONVERSION)+3, column=1,
         value='يمين = x تزيد | يسار = x تقل | فوق = y تقل | تحت = y تزيد').font = \
    Font(name='Calibri', bold=True, size=11, color='CC0000')
ws2.merge_cells(start_row=len(CONVERSION)+3, start_column=1,
                end_row=len(CONVERSION)+3, end_column=3)
ws2.cell(row=len(CONVERSION)+3, column=1).alignment = \
    Alignment(horizontal='center', readingOrder=2)

# ══════════════════════════════════════════════════════════════
# حفظ
# ══════════════════════════════════════════════════════════════
out = '/home/user/odoo_erp_modules_AR_default/FORM50_FIELD_MAP.xlsx'
wb.save(out)
print(f'✅ تم الحفظ: {out}')
